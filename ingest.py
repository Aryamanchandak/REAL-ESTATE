"""
ingest.py — Parses IPRU_RE_Fund_DataEntry_ver_3_final.xlsx into a dict of clean DataFrames.

Usage:
    from ingest import ingest
    tables = ingest("IPRU_RE_Fund_DataEntry_ver_3_final.xlsx")

Returns named DataFrames:
    meta                                    — fund name, report month, 6-month label array
    assets                                  — one row per physical asset
    portfolio_kpis                          — portfolio-level KPIs (key-value)
    portfolio_rent_spread                   — leased area by rent band
    portfolio_revenue_mix                   — revenue breakdown by line
    portfolio_asset_overview                — asset-level leasable/leased/vacant/occupancy (if present)
    portfolio_rent_spread_by_asset          — rent-band leased area per asset (if present)
    portfolio_revenue_mix_by_asset          — revenue-line breakdown per asset (if present)
    exec_asset_overview                     — per-asset deep metrics
    exec_rent_billed_vs_actual              — billed vs actual per asset × month
    exec_psf_trend                          — rent ₹/sf per asset × quarter
    exec_top_tenants                        — top tenants per asset
    financial_cash_position                 — opening/inflow/outflow/closing per asset
    financial_sources                       — sources of funds per asset
    financial_usage                         — usage of funds per asset
    financial_bp_vs_actual                  — business plan vs actual per asset × month
    financial_bp_vs_actual_total            — business plan vs actual, portfolio total per month
    financial_cash_flow_by_asset            — cash flow per asset
    expenses_monthly                        — expense vs collection per asset × month
    expenses_monthly_total                  — expense vs collection, portfolio total per month
    expenses_rent_billed_vs_actual          — billed vs actual per asset × month
    expenses_rent_billed_vs_actual_total    — billed vs actual, portfolio total per month
    stack_plan                              — floor-by-floor stacking plan
    tenant_profile                          — tenant list with rent and LED
    tenant_lease_expiry                     — lease expiry schedule by quarter
    approvals                               — statutory approval register
    compliance                              — compliance register
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

# ── constants ─────────────────────────────────────────────────────────────────

# First-column values that mark aggregation/check rows — always skipped
_SKIP_LABELS = {
    "TOTAL",
    "Total",
    "Grand Total",
    "Check (Open + In - Out = Close)",
}

# Null placeholder used throughout the template
_NULL_STR = "—"


# ── cell-level cleaning ───────────────────────────────────────────────────────

def _clean(v):
    """Normalise one cell value."""
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return None if s in (_NULL_STR, "") else s
    if isinstance(v, float):
        # Eliminate floating-point noise (e.g. 0.6400000000000006 → 0.64)
        return round(v, 6)
    return v


# ── row-level filters ─────────────────────────────────────────────────────────

def _is_skip_row(row: list) -> bool:
    """True for rows that should never appear in output DataFrames."""
    first = row[0]
    if not isinstance(first, str):
        return False
    return (
        first in _SKIP_LABELS
        or first.startswith("➕")
        or first.startswith("Category codes")
        or bool(re.match(r"^[A-Z]\d* · ", first))        # section headers (e.g. "A · ", "B1 · ")
        or first.startswith("STACKING PLAN")
        or first.startswith("APPROVAL REGISTER")
        or first.startswith("COMPLIANCE REGISTER")
        or first.startswith("ASSETS  —")
        or first.startswith("REPORT METADATA")
    )


# ── DataFrame builder ─────────────────────────────────────────────────────────

def _build_df(raw_rows: list, headers: list[str]) -> pd.DataFrame:
    """
    Convert a slice of raw openpyxl rows + a header list into a clean DataFrame.
    Skips fully-blank rows and skip-label rows.
    """
    records = []
    for row in raw_rows:
        row = list(row)
        if all(v is None for v in row):
            continue
        if _clean(row[0]) is None:      # no key value → filler/placeholder row
            continue
        if _is_skip_row(row):
            continue
        records.append([_clean(v) for v in row])
    df = pd.DataFrame(records, columns=headers).reset_index(drop=True)
    # Drop phantom columns: unnamed sheet columns beyond the section's real width
    phantom = [c for c in df.columns if c.startswith("_col")]
    return df.drop(columns=phantom)


def _headers(row) -> list[str]:
    return [str(h).strip() if h is not None else f"_col{i}" for i, h in enumerate(row)]


# ── section-anchor machinery ──────────────────────────────────────────────────

def _find_sections(ws) -> dict[str, int]:
    """
    Scan a worksheet and return {anchor: 1-based-row-number} for every row
    whose first cell matches the pattern 'X · <text>' or 'X1 · <text>'
    (e.g. 'A · ASSET OVERVIEW', 'D1 · MONTHLY TOTALS').

    Anchors are keyed by their full token ("A", "B1", ...) — not just the
    leading letter — since a sheet can contain both "B" and "B1" as distinct
    sections.
    """
    out = {}
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        v = row[0]
        if isinstance(v, str):
            m = re.match(r"^([A-Z]\d*) · ", v)
            if m:
                out[m.group(1)] = i
    return out


def _ordered_anchors(sections: dict[str, int]) -> list[tuple[str, int]]:
    """
    Anchor tokens sorted by physical row position, not alphabetically.

    Sub-anchors like "B1"/"D1" sort after their parent letter alphabetically
    ("A" < "B" < "B1" < "C") but are not necessarily adjacent in the sheet
    (e.g. PORTFOLIO's physical order is A, B, C, D, B1, C1). Section
    boundaries must follow physical order or slicing silently swallows the
    wrong rows.
    """
    return sorted(sections.items(), key=lambda kv: kv[1])


def _read_section(
    all_rows: list,
    ordered: list[tuple[str, int]],
    anchor: str,
    max_row: int,
) -> pd.DataFrame:
    """
    Read one anchored section from a pre-loaded list of rows.

    Layout assumption (consistent across all multi-section sheets):
        row[s-1]  (0-based) = section header  "X · ..."   ← skipped
        row[s]    (0-based) = column headers               ← used as headers
        row[s+1:] (0-based) = data rows                    ← parsed to DataFrame

    The section's end boundary is the next anchor in physical (row) order,
    regardless of alphabetical relationship to 'anchor'.
    """
    idx = next(i for i, (tok, _) in enumerate(ordered) if tok == anchor)
    s = ordered[idx][1]   # 1-based row of this anchor's title row
    col_hdr_idx = s       # 0-based index of the column-header row
    data_start   = s + 1  # 0-based index of first data row

    if idx + 1 < len(ordered):
        data_end = ordered[idx + 1][1] - 1   # exclusive, 0-based
    else:
        data_end = max_row  # ws.max_row is 1-based; used as exclusive end → reads all remaining

    hdrs = _headers(all_rows[col_hdr_idx])
    return _build_df(all_rows[data_start:data_end], hdrs)


# ── per-sheet ingesters ───────────────────────────────────────────────────────

def _ingest_meta(ws) -> pd.DataFrame:
    """
    Sheet: 1 · META
    Structure: row 1 = title, row 2 = headers, row 3 = single data row (9 cols).
    Cols D–I (index 3–8) are the 6-month label array.
    """
    rows = list(ws.iter_rows(values_only=True))
    hdrs = _headers(rows[1])
    data = [_clean(v) for v in rows[2]]
    return pd.DataFrame([data], columns=hdrs)


def _ingest_assets(ws) -> pd.DataFrame:
    """
    Sheet: 2 · ASSETS
    Structure: row 1 = title, row 2 = headers, rows 3–N = one row per asset
               (variable count; TOTAL row skipped by _SKIP_LABELS).
    """
    rows = list(ws.iter_rows(values_only=True))
    hdrs = _headers(rows[1])
    return _build_df(rows[2:], hdrs)


def _ingest_portfolio(ws) -> dict[str, pd.DataFrame]:
    """
    Sheet: 3 · PORTFOLIO
    A · PORTFOLIO KPIs         → 8 KPI rows
    B · RENT SPREAD            → 4 rent-band rows (Total filtered)
    C · REVENUE COMPOSITION    → 4 revenue lines (Total filtered)
    D · ASSET-LEVEL OVERVIEW   → one row per asset (Total filtered)
    B1 · RENT SPREAD BY ASSET  → one row per asset (Total filtered)
    C1 · REVENUE COMP BY ASSET → one row per asset (Total filtered)

    Physical sheet order is A, B, C, D, B1, C1 — NOT alphabetical — so
    boundaries must come from _ordered_anchors, not sorted(ss).
    """
    rows = list(ws.iter_rows(values_only=True))
    ss   = _find_sections(ws)
    mr   = ws.max_row
    ordered = _ordered_anchors(ss)
    out = {
        "portfolio_kpis":               _read_section(rows, ordered, "A", mr),
        "portfolio_rent_spread":        _read_section(rows, ordered, "B", mr),
        "portfolio_revenue_mix":        _read_section(rows, ordered, "C", mr),
    }
    if "D" in ss:
        out["portfolio_asset_overview"] = _read_section(rows, ordered, "D", mr)
    if "B1" in ss:
        out["portfolio_rent_spread_by_asset"] = _read_section(rows, ordered, "B1", mr)
    if "C1" in ss:
        out["portfolio_revenue_mix_by_asset"] = _read_section(rows, ordered, "C1", mr)
    return out


def _ingest_executive(ws) -> dict[str, pd.DataFrame]:
    """
    Sheet: 4 · EXECUTIVE  (4 sections)
    Row counts scale with the number of assets/months/tenants in the workbook
    (all sections are parsed dynamically, not assumed to be a fixed size):
    A · ASSET OVERVIEW         → one row per asset
    B · RENT BILLED vs ACTUAL  → assets × months
    C · RENTAL PSF TREND       → assets × quarters
    D · TOP TENANTS            → assets × tenants
    """
    rows = list(ws.iter_rows(values_only=True))
    ss   = _find_sections(ws)
    mr   = ws.max_row
    ordered = _ordered_anchors(ss)
    names = {
        "A": "exec_asset_overview",
        "B": "exec_rent_billed_vs_actual",
        "C": "exec_psf_trend",
        "D": "exec_top_tenants",
    }
    out = {}
    for anchor, _ in ordered:
        if anchor in names:
            out[names[anchor]] = _read_section(rows, ordered, anchor, mr)
    return out


def _ingest_financial(ws) -> dict[str, pd.DataFrame]:
    """
    Sheet: 5 · FINANCIAL
    A · CASH POSITION           → one row per asset (Check filtered)
    B · SOURCES OF FUNDS        → one row per asset (Grand Total filtered)
    C · USAGE OF FUNDS          → one row per asset (Grand Total filtered)
    D · BUSINESS PLAN vs ACTUAL → one row per asset × month (Total filtered)
    D1 · MONTHLY TOTALS         → portfolio-level, one row per month
    E · CASH FLOW BY ASSET      → one row per asset (TOTAL filtered)

    Physical order is A, B, C, D, D1, E.
    """
    rows = list(ws.iter_rows(values_only=True))
    ss   = _find_sections(ws)
    mr   = ws.max_row
    ordered = _ordered_anchors(ss)
    names = {
        "A":  "financial_cash_position",
        "B":  "financial_sources",
        "C":  "financial_usage",
        "D":  "financial_bp_vs_actual",
        "D1": "financial_bp_vs_actual_total",
        "E":  "financial_cash_flow_by_asset",
    }
    out = {}
    for anchor, _ in ordered:
        if anchor in names:
            out[names[anchor]] = _read_section(rows, ordered, anchor, mr)
    return out


def _ingest_expenses(ws) -> dict[str, pd.DataFrame]:
    """
    Sheet: 6 · EXPENSES
    A · MONTHLY EXPENSE vs COLLECTION → one row per asset × month
    A1 · MONTHLY TOTALS               → portfolio-level, one row per month
    B · RENT BILLED vs ACTUAL         → one row per asset × month
    B1 · MONTHLY TOTALS               → portfolio-level, one row per month

    Physical order is A, A1, B, B1.
    """
    rows = list(ws.iter_rows(values_only=True))
    ss   = _find_sections(ws)
    mr   = ws.max_row
    ordered = _ordered_anchors(ss)
    names = {
        "A":  "expenses_monthly",
        "A1": "expenses_monthly_total",
        "B":  "expenses_rent_billed_vs_actual",
        "B1": "expenses_rent_billed_vs_actual_total",
    }
    out = {}
    for anchor, _ in ordered:
        if anchor in names:
            out[names[anchor]] = _read_section(rows, ordered, anchor, mr)
    return out


def _ingest_stack_plan(ws) -> pd.DataFrame:
    """
    Sheet: 7 · STACK PLAN
    Structure: rows 1–2 = preamble (skipped), row 3 = headers, rows 4–N = data.
    Variable-length: users can append rows above the ➕ sentinel.
    Sentinel row and fully-blank rows are filtered.
    """
    rows = list(ws.iter_rows(values_only=True))
    hdrs = _headers(rows[2])          # row 3 (0-based index 2)
    return _build_df(rows[3:], hdrs)  # row 4 onward


def _ingest_tenant(ws) -> dict[str, pd.DataFrame]:
    """
    Sheet: 8 · TENANT  (2 sections)
    A · TENANT PROFILE       → variable rows per asset×tenant
    B · LEASE EXPIRY SCHEDULE→ variable rows per asset×quarter
    Note: Section B has no column-header row; headers are hardcoded here.
    """
    rows = list(ws.iter_rows(values_only=True))
    ss   = _find_sections(ws)
    mr   = ws.max_row
    ordered = _ordered_anchors(ss)
    # Section B data starts immediately after the section marker (no header row)
    b = ss["B"]  # 1-based row of the marker; 0-based index b = first data row
    lease_hdrs = ["Asset ID", "Quarter", "Area Expiring (sft)", "Note"]
    return {
        "tenant_profile":      _read_section(rows, ordered, "A", mr),
        "tenant_lease_expiry": _build_df(rows[b:mr], lease_hdrs),
    }


def _ingest_approvals(ws) -> pd.DataFrame:
    """
    Sheet: 9 · APPROVALS
    Structure: row 1 = title+status legend, row 2 = headers, rows 3–N = data.
    """
    rows = list(ws.iter_rows(values_only=True))
    hdrs = _headers(rows[1])
    return _build_df(rows[2:], hdrs)


def _ingest_compliance(ws) -> pd.DataFrame:
    """
    Sheet: 10 · COMPLIANCE
    Structure: row 1 = title+RAG legend, row 2 = headers, rows 3–N = data.
    """
    rows = list(ws.iter_rows(values_only=True))
    hdrs = _headers(rows[1])
    return _build_df(rows[2:], hdrs)


# ── public entry point ────────────────────────────────────────────────────────

def ingest(filepath: str | Path) -> dict[str, pd.DataFrame]:
    """
    Parse the DataEntry workbook and return a flat dict of DataFrames.
    Sheet names and column order must match the template exactly.
    """
    wb = openpyxl.load_workbook(str(filepath), data_only=True)

    tables: dict[str, pd.DataFrame] = {}
    tables["meta"]   = _ingest_meta(wb["1 · META"])
    tables["assets"] = _ingest_assets(wb["2 · ASSETS"])
    tables.update(_ingest_portfolio(wb["3 · PORTFOLIO"]))
    tables.update(_ingest_executive(wb["4 · EXECUTIVE"]))
    tables.update(_ingest_financial(wb["5 · FINANCIAL"]))
    tables.update(_ingest_expenses(wb["6 · EXPENSES"]))
    tables["stack_plan"]  = _ingest_stack_plan(wb["7 · STACK PLAN"])
    tables.update(_ingest_tenant(wb["8 · TENANT"]))
    tables["approvals"]  = _ingest_approvals(wb["9 · APPROVALS"])
    tables["compliance"] = _ingest_compliance(wb["10 · COMPLIANCE"])

    return tables


# ── CLI smoke-test ────────────────────────────────────────────────────────────

if __name__ == "__main__":
    import sys

    path = sys.argv[1] if len(sys.argv) > 1 else "IPRU_RE_Fund_DataEntry_ver_3_final.xlsx"
    tables = ingest(path)

    for name, df in tables.items():
        print(f"\n{'─' * 60}")
        print(f"  {name}  ({len(df)} rows × {len(df.columns)} cols)")
        print(f"{'─' * 60}")
        print(df.to_string(index=False))
