"""
ingest.py — reads the IPRU Real Estate Fund workbook (simplified, 9-asset-ready format).

Usage:
    from ingest import ingest
    book = ingest("data/real_estate_fund_data.xlsx")

Sources (sheet → what is read):
    0_README                 fund identity (Fund Name, Currency, Reporting Unit …)
    DATA_QUALITY_CHECKS      the workbook's own publishing checklist
    1_ASSET_MASTER           asset identity, address, image, map, capex till date
    2_MONTHLY_INPUT          one row per asset × month — every monthly figure
    3_REVENUE_RENT           rent spread: FUND table (sft) + by-asset matrix (M sft)
    3A_REVENUE_MIX           optional revenue lines per month × asset
    4_TENANT_STACK_LEASE     three stacked blocks: tenant profile, lease expiry, stack plan
    6_APPROVALS_COMPLIANCE   two side-by-side registers
    7_FUND_MONTHLY_SUMMARY   read only to cross-check the fund roll-up

The export/summary sheets (5_DASHBOARD_EXPORT, 10_TECH_EXPORT_*, 8/9 summaries) are
formula mirrors of 2_MONTHLY_INPUT that turn blanks into zeros and drop columns, so
fund and annual figures are rolled up here from the input rows instead.

Unit conventions (applied once, here):
    money      ₹ Crore, as entered
    area       square feet; stack-plan areas stay in K sft, as entered
    occupancy  percent, 0–100
    months     datetime.date on the first of the month
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import quote_plus
import warnings

import openpyxl
from openpyxl.utils.datetime import from_excel

FUND_ID = "FUND"

SHEET_README = "0_README"
SHEET_QUALITY = "DATA_QUALITY_CHECKS"
SHEET_ASSETS = "1_ASSET_MASTER"
SHEET_MONTHLY = "2_MONTHLY_INPUT"
SHEET_RENT = "3_REVENUE_RENT"
SHEET_REVENUE_MIX = "3A_REVENUE_MIX"
SHEET_TENANT = "4_TENANT_STACK_LEASE"
SHEET_REGISTERS = "6_APPROVALS_COMPLIANCE"
SHEET_FUND_MONTHLY = "7_FUND_MONTHLY_SUMMARY"

REQUIRED_SHEETS = (SHEET_ASSETS, SHEET_MONTHLY, SHEET_RENT, SHEET_TENANT, SHEET_REGISTERS)

TOLERANCE_CR = 0.005

# key → canonical 2_MONTHLY_INPUT header
MONTHLY_FIELDS = {
    "leasable": "Leasable Area (sft)",
    "leased": "Leased Area (sft)",
    "psf": "Current Weighted Average PSF Rent (INR)",
    "monthlyRent": "Monthly Rent (₹Cr)",
    "monthlyCam": "Monthly CAM (₹Cr)",
    "grossBilling": "Gross Billing (₹Cr)",
    "rentalCollections": "Rental Collections (₹Cr)",
    "camCollections": "CAM Collections (₹Cr)",
    "otherCollections": "Other Collections (₹Cr)",
    "collections": "Collections Inflow (₹Cr)",
    "openingBalance": "Opening Balance (₹Cr)",
    "overdraftOpening": "Over Draft Opening Balance (₹Cr)",
    "redemptionFd": "Redemption of FD (₹Cr)",
    "otherInflow": "Other Inflow (₹Cr)",
    "totalInflow": "Total Inflow (₹Cr)",
    "statutory": "Statutory Payments (TDS/GST) (₹Cr)",
    "camExpense": "CAM Expense (₹Cr)",
    "debtService": "Debt Service (₹Cr)",
    "opex": "OPEX (₹Cr)",
    "totalExpense": "Total Expense (₹Cr)",
    "netSurplus": "Net Surplus (₹Cr)",
    "investmentMf": "Investment in MF (₹Cr)",
    "creationFd": "Creation of FD (₹Cr)",
    "capex": "Capex (₹Cr)",
    "otherOutflow": "Other Outflow (₹Cr)",
    "bankCharges": "Bank Charges (₹Cr)",
    "totalOutflow": "Total Outflow (₹Cr)",
    "closingBalance": "Closing Balance (₹Cr)",
    "rentToBeBilled": "Rent to be Billed (₹Cr)",
    "actualBilled": "Actual Billed (₹Cr)",
    "plannedCollection": "Planned Collection / BP (₹Cr)",
    "tdsReceivable": "TDS Receivable (₹Cr)",
}
OPTIONAL_MONTHLY = {
    "overdraftOpening", "investmentMf", "tdsReceivable",
    "rentToBeBilled", "actualBilled", "plannedCollection",
}

# A month counts as reported only when one of these hand-entered fields has a value.
# Opening Balance is excluded: it rolls forward by formula into months not yet filled.
EVIDENCE_FIELDS = (
    "leased", "psf", "monthlyRent", "monthlyCam", "rentalCollections", "camCollections",
    "otherCollections", "redemptionFd", "otherInflow", "statutory", "camExpense",
    "debtService", "opex", "investmentMf", "creationFd", "capex", "otherOutflow",
    "bankCharges", "rentToBeBilled", "actualBilled", "plannedCollection",
)

COLLECTION_PARTS = ("rentalCollections", "camCollections", "otherCollections")
INFLOW_PARTS = ("collections", "redemptionFd", "otherInflow")
EXPENSE_PARTS = ("statutory", "camExpense", "debtService", "opex")
OUTFLOW_PARTS = ("totalExpense", "investmentMf", "creationFd", "capex", "otherOutflow", "bankCharges")

# Summed across assets (fund) and across months (annual).
FLOW_FIELDS = (
    "monthlyRent", "monthlyCam", "grossBilling", *COLLECTION_PARTS, "collections",
    "redemptionFd", "otherInflow", "totalInflow", *EXPENSE_PARTS, "totalExpense",
    "netSurplus", "investmentMf", "creationFd", "capex", "otherOutflow", "bankCharges",
    "totalOutflow", "outflowComponents", "rentToBeBilled", "actualBilled",
    "plannedCollection",
)
# Summed across assets only (point-in-time values).
STOCK_FIELDS = (
    "leasable", "leased", "vacant", "openingBalance", "closingBalance", "tdsReceivable",
)

_NULL_STRINGS = {"", "—", "–"}


@dataclass
class IngestionReport:
    active_asset_ids: list = field(default_factory=list)
    dropped_rows: dict = field(default_factory=dict)
    warnings: list = field(default_factory=list)

    def warn(self, message: str) -> None:
        if message not in self.warnings:
            self.warnings.append(message)

    def dropped(self, table: str, count: int = 1) -> None:
        self.dropped_rows[table] = self.dropped_rows.get(table, 0) + count


@dataclass
class FundWorkbook:
    meta: dict
    assets: list
    months: list
    report_month: date
    as_of: date
    monthly: dict
    annual: dict
    rent_spread: dict
    revenue_mix: list
    tenants: list
    lease_schedule: list
    stack_plan: list
    approvals: list
    compliance: list
    data_quality: dict
    report: IngestionReport


# ── cell helpers ──────────────────────────────────────────────────────────────

def _canon(header):
    if header is None:
        return None
    text = re.sub(r"\s+", " ", str(header)).strip()
    text = re.sub(r"\(\s+", "(", text)
    return re.sub(r"\s+\)", ")", text) or None


def _clean(value):
    if isinstance(value, str):
        text = re.sub(r"\s+", " ", value).strip()
        return None if text in _NULL_STRINGS else text
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def _num(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return None if math.isnan(float(value)) else float(value)
    if isinstance(value, str):
        try:
            return float(value.replace(",", ""))
        except ValueError:
            return None
    return None


def _to_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and 20000 < value < 80000:
        return from_excel(value).date()
    if isinstance(value, str):
        for fmt in ("%d %b %Y", "%d-%b-%Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _to_month(value):
    if isinstance(value, str):
        for fmt in ("%b-%y", "%b-%Y", "%Y-%m", "%B %Y", "%b %Y"):
            try:
                parsed = datetime.strptime(value.strip(), fmt)
                return date(parsed.year, parsed.month, 1)
            except ValueError:
                continue
        return None
    parsed = _to_date(value)
    return date(parsed.year, parsed.month, 1) if parsed else None


def _sum(values):
    present = [v for v in values if v is not None]
    return sum(present) if present else None


def _differs(a, b, tolerance=TOLERANCE_CR):
    return a is not None and b is not None and abs(a - b) > tolerance


def annual_period(month: date) -> str:
    """June–May reporting year label, matching the workbook ('Jun-26 to May-27')."""
    start = month.year if month.month >= 6 else month.year - 1
    return f"Jun-{start % 100:02d} to May-{(start + 1) % 100:02d}"


def month_label(month: date) -> str:
    return month.strftime("%b-%y")


def quarter_label(day: date) -> str:
    return f"{day.year % 100:02d} Q{(day.month - 1) // 3 + 1}"


def _month_end(month: date) -> date:
    following = date(month.year + (month.month == 12), month.month % 12 + 1, 1)
    return following - timedelta(days=1)


def _safe_url(value, report: IngestionReport, where: str):
    if not isinstance(value, str) or not value.strip():
        return None
    url = value.strip()
    if re.match(r"^https?://", url, re.I):
        return url
    report.warn(f"{where}: ignored non-http(s) link {url[:60]!r}")
    return None


# ── sheet slicing ─────────────────────────────────────────────────────────────

def _rows(ws) -> list:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _find_row(rows, predicate, start=0):
    for index in range(start, len(rows)):
        if predicate(rows[index]):
            return index
    return None


def _first_cell_is(text: str):
    target = text.lower()
    return lambda row: bool(row) and isinstance(row[0], str) and _canon(row[0]).lower() == target


def _block(rows, header_index, end_index=None, first_col=0, last_col=None) -> tuple:
    """(headers, records) for a table whose header row is rows[header_index]."""
    headers = [_canon(h) for h in rows[header_index][first_col:last_col]]
    records = []
    for row in rows[header_index + 1:end_index]:
        cells = [_clean(v) for v in row[first_col:last_col]]
        cells += [None] * (len(headers) - len(cells))
        if all(c is None for c in cells):
            continue
        records.append({h: c for h, c in zip(headers, cells) if h})
    return [h for h in headers if h], records


def _require(headers, required, where):
    missing = [column for column in required if column not in headers]
    if missing:
        raise ValueError(f"{where}: missing required column(s): {missing}")


def _header_row(rows, first_header, where):
    index = _find_row(rows, _first_cell_is(first_header))
    if index is None:
        raise ValueError(f"{where}: header row starting with {first_header!r} not found")
    return index


# ── readers ───────────────────────────────────────────────────────────────────

def _read_meta(wb) -> dict:
    meta = {"fund": None, "currency": None, "unit": None, "frequency": None, "annualConvention": None}
    if SHEET_README not in wb.sheetnames:
        return meta
    labels = {
        "fund name": "fund", "currency": "currency", "reporting unit": "unit",
        "reporting frequency": "frequency", "annual period convention": "annualConvention",
    }
    for row in _rows(wb[SHEET_README]):
        key = _canon(row[0]) if row and isinstance(row[0], str) else None
        if key and key.lower() in labels and len(row) > 1:
            meta[labels[key.lower()]] = _clean(row[1])
    return meta


def _read_data_quality(wb) -> dict:
    result = {"overall": None, "checks": []}
    if SHEET_QUALITY not in wb.sheetnames:
        return result
    rows = _rows(wb[SHEET_QUALITY])
    index = _find_row(rows, _first_cell_is("Check"))
    if index is None:
        return result
    _, records = _block(rows, index)
    for record in records:
        check = record.get("Check")
        if not check:
            continue
        if check.lower().startswith("overall"):
            result["overall"] = record.get("Status")
            continue
        result["checks"].append({
            "check": check,
            "status": record.get("Status"),
            "outstanding": record.get("Outstanding Records"),
            "action": record.get("Action"),
        })
    return result


def _read_assets(wb, report) -> list:
    rows = _rows(wb[SHEET_ASSETS])
    index = _header_row(rows, "Asset ID", SHEET_ASSETS)
    headers, records = _block(rows, index)
    _require(headers, ("Asset ID", "Asset Name", "City, State", "Address", "Leasable (sft)"), SHEET_ASSETS)

    assets, seen = [], set()
    for record in records:
        asset_id, name = record.get("Asset ID"), record.get("Asset Name")
        if asset_id is None or name is None:
            continue   # unused capacity slot
        asset_id = str(asset_id)
        if asset_id == FUND_ID:
            raise ValueError(f"{SHEET_ASSETS}: {FUND_ID!r} is reserved for the fund roll-up")
        if asset_id in seen:
            raise ValueError(f"{SHEET_ASSETS}: duplicate Asset ID {asset_id!r}")
        seen.add(asset_id)

        where = f"{SHEET_ASSETS}[{asset_id}]"
        lat, lng = _num(record.get("Latitude")), _num(record.get("Longitude"))
        address = record.get("Address") or f"{name}, {record.get('City, State') or ''}".strip(", ")
        embed = _safe_url(record.get("Google Maps Embed URL / Iframe Src"), report, where)
        place_link = embed
        # Google only allows /maps/embed or output=embed URLs inside an iframe; a
        # pasted /maps/place/ link is kept as the "open in Maps" link instead.
        if embed and "output=embed" not in embed and "/maps/embed" not in embed:
            embed = None
        if embed is None and lat is not None and lng is not None:
            embed = f"https://maps.google.com/maps?q={lat},{lng}&z=16&output=embed"
        search = "https://www.google.com/maps/search/?api=1&query=" + quote_plus(address)

        assets.append({
            "id": asset_id,
            "name": name,
            "city": record.get("City, State"),
            "address": record.get("Address"),
            "leasable": _num(record.get("Leasable (sft)")),
            "imageUrl": _safe_url(record.get("Asset Image URL / File Link"), report, where),
            "imageAlt": record.get("Asset Image Alt Text") or name,
            "mapEmbedUrl": embed,
            "mapLinkUrl": place_link or search,
            "latitude": lat,
            "longitude": lng,
            "notes": record.get("Notes"),
            "capexTillDate": _num(record.get("Total Capex Till Date (₹Cr)")),
            "workbookWale": _num(record.get("Weighted Average Lease Expiry (yrs)")),
        })
    if not assets:
        raise ValueError(f"{SHEET_ASSETS}: no active assets (Asset ID + Asset Name) found")
    report.active_asset_ids = [a["id"] for a in assets]
    return assets


def _read_monthly(wb, assets, report) -> tuple:
    rows = _rows(wb[SHEET_MONTHLY])
    index = _header_row(rows, "Month", SHEET_MONTHLY)
    headers, records = _block(rows, index)
    required = [h for k, h in MONTHLY_FIELDS.items() if k not in OPTIONAL_MONTHLY]
    _require(headers, ["Month", "Asset ID", *required], SHEET_MONTHLY)

    active = {a["id"] for a in assets}
    parsed, unknown = [], set()
    for record in records:
        asset_id, month = record.get("Asset ID"), _to_month(record.get("Month"))
        if asset_id is None or month is None:
            continue
        asset_id = str(asset_id)
        row = {"assetId": asset_id, "month": month, "note": record.get("Note")}
        for key, header in MONTHLY_FIELDS.items():
            raw = record.get(header)
            row[key] = _num(raw)
            if raw is not None and row[key] is None:
                report.warn(f"{SHEET_MONTHLY}[{asset_id} {month_label(month)}]: {header} is not numeric: {raw!r}")
        has_evidence = any(row[k] is not None for k in EVIDENCE_FIELDS)
        if asset_id not in active:
            if has_evidence:
                unknown.add(asset_id)
                report.dropped(SHEET_MONTHLY)
            continue
        row["hasData"] = has_evidence
        parsed.append(row)
    if unknown:
        report.warn(f"{SHEET_MONTHLY}: dropped rows for Asset ID(s) not in {SHEET_ASSETS}: {sorted(unknown)}")

    reported = sorted({r["month"] for r in parsed if r["hasData"]})
    if not reported:
        raise ValueError(f"{SHEET_MONTHLY}: no month has any entered figures")
    report_month = reported[-1]
    gaps = [m for m in {r["month"] for r in parsed} if reported[0] < m < report_month and m not in reported]
    if gaps:
        report.warn(f"{SHEET_MONTHLY}: no figures entered for {', '.join(month_label(m) for m in sorted(gaps))}")

    by_asset = {a["id"]: [] for a in assets}
    for row in sorted(parsed, key=lambda r: r["month"]):
        if row["month"] in reported:
            _derive_monthly_row(row, report)
            by_asset[row["assetId"]].append(row)

    for asset_id, series in by_asset.items():
        if not any(r["hasData"] for r in series):
            report.warn(f"{SHEET_MONTHLY}: no monthly figures entered for asset {asset_id!r}")
        latest = [r for r in series if r["month"] == report_month]
        if latest and not latest[0]["hasData"]:
            report.warn(f"{SHEET_MONTHLY}: {asset_id!r} has no figures for the report month {month_label(report_month)}")
        for previous, current in zip(series, series[1:]):
            if _differs(previous["closingBalance"], current["openingBalance"]):
                report.warn(
                    f"{SHEET_MONTHLY}[{asset_id} {month_label(current['month'])}]: opening balance "
                    f"{current['openingBalance']:.4f} differs from prior month closing "
                    f"{previous['closingBalance']:.4f}"
                )
    return by_asset, reported, report_month


def _derive_monthly_row(row, report):
    where = f"{SHEET_MONTHLY}[{row['assetId']} {month_label(row['month'])}]"

    # The sheet computes these; recompute from their inputs so a hand-edited or
    # mis-referenced formula (seen in Net Surplus and Opening Balance) cannot leak through.
    if row["leasable"] is not None and row["leased"] is not None:
        row["vacant"] = row["leasable"] - row["leased"]
        row["occupancy"] = row["leased"] / row["leasable"] * 100 if row["leasable"] else None
    else:
        row["vacant"] = row["occupancy"] = None

    q, r = row["openingBalance"], row["overdraftOpening"]
    row["openingSource"] = "Opening Balance" if q is not None else ("Over Draft Opening Balance" if r is not None else None)
    if q is None:
        row["openingBalance"] = r
    elif _differs(q, r):
        report.warn(f"{where}: Opening Balance {q:.4f} and Over Draft Opening Balance {r:.4f} disagree; using Opening Balance")

    for total, parts, label in (
        ("collections", COLLECTION_PARTS, "Collections Inflow"),
        ("totalInflow", INFLOW_PARTS, "Total Inflow"),
        ("totalExpense", EXPENSE_PARTS, "Total Expense"),
    ):
        derived = _sum(row[p] for p in parts)
        if _differs(derived, row[total]):
            report.warn(f"{where}: workbook {label} {row[total]:.4f} ≠ sum of its components {derived:.4f}; using the sum")
        row[total] = derived

    net = row["collections"] - row["totalExpense"] if row["collections"] is not None and row["totalExpense"] is not None else None
    if _differs(net, row["netSurplus"]):
        report.warn(
            f"{where}: workbook Net Surplus {row['netSurplus']:.4f} is not Collections Inflow − Total Expense "
            f"({net:.4f}); showing {net:.4f}"
        )
    row["netSurplus"] = net

    # Total Outflow and Closing Balance are reported balances, so they are kept as
    # entered; the dashboard shows any outflow component the workbook left out.
    row["outflowComponents"] = _sum(row[p] for p in OUTFLOW_PARTS)
    gap = (row["outflowComponents"] or 0) - (row["totalOutflow"] or 0)
    row["outflowGap"] = gap if abs(gap) > TOLERANCE_CR else 0.0
    if row["outflowGap"]:
        report.warn(
            f"{where}: outflow components total {row['outflowComponents']:.4f} but workbook Total Outflow is "
            f"{(row['totalOutflow'] or 0):.4f} (gap {gap:.4f}"
            + (f"; Investment in MF {row['investmentMf']:.4f} is not in Total Outflow" if row["investmentMf"] else "")
            + ") — Closing Balance follows the workbook"
        )

    if row["hasData"] and row["closingBalance"] is None and row["openingBalance"] is not None:
        report.warn(f"{where}: Closing Balance has no saved value; open, recalculate and save the workbook in Excel")
    if None not in (row["openingBalance"], row["totalInflow"], row["totalOutflow"], row["closingBalance"]):
        delta = row["openingBalance"] + row["totalInflow"] - row["totalOutflow"] - row["closingBalance"]
        if abs(delta) > TOLERANCE_CR:
            report.warn(f"{where}: opening + inflow − outflow differs from closing by {delta:.4f}")


def _roll_up_fund(by_asset, months) -> list:
    fund = []
    for month in months:
        rows = [r for series in by_asset.values() for r in series if r["month"] == month]
        out = {"assetId": FUND_ID, "month": month, "hasData": any(r["hasData"] for r in rows), "note": None}
        for key in FLOW_FIELDS + STOCK_FIELDS + ("outflowGap",):
            out[key] = _sum(r.get(key) for r in rows)
        with_leased = [r for r in rows if r["leased"] is not None and r["leasable"]]
        leasable = sum(r["leasable"] for r in with_leased)
        out["occupancy"] = sum(r["leased"] for r in with_leased) / leasable * 100 if leasable else None
        priced = [r for r in rows if r["psf"] is not None and r["leased"]]
        leased = sum(r["leased"] for r in priced)
        out["psf"] = sum(r["psf"] * r["leased"] for r in priced) / leased if leased else None
        out["psfExcluded"] = [r["assetId"] for r in rows if r["leased"] and r["psf"] is None]
        out["openingSource"] = None
        out["overdraftOpening"] = _sum(r["overdraftOpening"] for r in rows)
        fund.append(out)
    return fund


def _check_fund_summary(wb, fund, report):
    """Compare the Python roll-up with 7_FUND_MONTHLY_SUMMARY and explain differences."""
    if SHEET_FUND_MONTHLY not in wb.sheetnames:
        return
    rows = _rows(wb[SHEET_FUND_MONTHLY])
    index = _find_row(rows, _first_cell_is("Month"))
    if index is None:
        return
    _, records = _block(rows, index)
    sheet = {_to_month(r.get("Month")): r for r in records if _to_month(r.get("Month"))}
    checks = (
        ("leasable", "Total Leasable Area", 1), ("leased", "Total Leased Area", 1),
        ("monthlyRent", "Monthly Rent", TOLERANCE_CR), ("grossBilling", "Gross Billing", TOLERANCE_CR),
        ("collections", "Collections Inflow", TOLERANCE_CR), ("totalExpense", "Total Expense", TOLERANCE_CR),
        ("openingBalance", "Opening Balance", TOLERANCE_CR), ("closingBalance", "Closing Balance", TOLERANCE_CR),
        ("psf", "Current Weighted Average PSF Rent (INR)", 0.05),
    )
    for row in fund:
        record = sheet.get(row["month"])
        if not record:
            continue
        for key, header, tolerance in checks:
            workbook_value = _num(record.get(header))
            if _differs(row[key], workbook_value, tolerance):
                reason = ""
                if key == "psf" and row["psfExcluded"]:
                    reason = f" (the sheet counts blank PSF as 0 for {', '.join(row['psfExcluded'])}; the dashboard leaves those assets out of the average)"
                elif key == "openingBalance":
                    reason = " (the sheet ignores Over Draft Opening Balance rows)"
                report.warn(
                    f"{SHEET_FUND_MONTHLY}[{month_label(row['month'])}]: {header} is {workbook_value:,.4f} in the sheet "
                    f"but {row[key]:,.4f} when rolled up from {SHEET_MONTHLY}{reason}"
                )


def _annual(series) -> list:
    periods = {}
    for row in series:
        periods.setdefault(annual_period(row["month"]), []).append(row)
    result = []
    for label, rows in periods.items():
        rows = sorted(rows, key=lambda r: r["month"])
        entry = {"period": label, "monthsReported": sum(1 for r in rows if r["hasData"]),
                 "firstMonth": rows[0]["month"], "lastMonth": rows[-1]["month"]}
        for key in FLOW_FIELDS + ("outflowGap",):
            entry[key] = _sum(r.get(key) for r in rows)
        for key in ("leasable", "leased", "occupancy", "psf"):
            values = [r[key] for r in rows if r.get(key) is not None]
            entry[key] = sum(values) / len(values) if values else None
        entry["openingBalance"] = rows[0]["openingBalance"]
        entry["closingBalance"] = rows[-1]["closingBalance"]
        result.append(entry)
    return result


def _read_rent_spread(wb, assets, report) -> dict:
    rows = _rows(wb[SHEET_RENT])
    index = _header_row(rows, "Rent Band", SHEET_RENT)
    headers, _ = _block(rows, index, index + 1)
    _require(headers, ("Rent Band", "Area (sft)"), SHEET_RENT)
    fund, as_of, end = [], None, None
    for offset, row in enumerate(rows[index + 1:], start=index + 1):
        record = dict(zip([_canon(h) for h in rows[index]], [_clean(v) for v in row]))
        band = record.get("Rent Band")
        if band is None or str(band).lower() == "total":
            end = offset
            break
        fund.append({"band": str(band), "sft": _num(record.get("Area (sft)")) or 0.0})
        as_of = as_of or _to_month(record.get("As Of Month"))

    # Below the FUND table sits a matrix: one row per Asset ID, one column per band, in M sft.
    bands = [b["band"] for b in fund]
    by_asset = {}
    matrix_header = _find_row(
        rows, lambda r: sum(1 for v in r[1:] if _clean(v) in bands) >= max(1, len(bands) // 2), (end or index) + 1
    )
    if matrix_header is not None:
        columns = {i: _clean(v) for i, v in enumerate(rows[matrix_header]) if _clean(v) in bands}
        active = {a["id"] for a in assets}
        for row in rows[matrix_header + 1:]:
            key = _clean(row[0])
            if key is None:
                continue
            if str(key).lower() == "total":
                break
            if str(key) not in active:
                report.warn(f"{SHEET_RENT}: rent-spread row for unknown Asset ID {key!r} ignored")
                continue
            by_asset[str(key)] = [
                {"band": band, "sft": (_num(row[i]) or 0.0) * 1_000_000 if i < len(row) else 0.0}
                for i, band in columns.items()
            ]
        for position, band in enumerate(bands):
            total = sum(values[position]["sft"] for values in by_asset.values() if position < len(values))
            if by_asset and abs(total - fund[position]["sft"]) > 1:
                report.warn(f"{SHEET_RENT}: {band} FUND area {fund[position]['sft']:,.0f} sft ≠ sum of assets {total:,.0f} sft")
    return {"fund": fund, "byAsset": by_asset, "asOf": as_of}


def _read_revenue_mix(wb, assets, report) -> list:
    if SHEET_REVENUE_MIX not in wb.sheetnames:
        return []
    rows = _rows(wb[SHEET_REVENUE_MIX])
    index = _find_row(rows, _first_cell_is("Month"))
    if index is None:
        return []
    _, records = _block(rows, index)
    active = {a["id"] for a in assets}
    result = []
    for record in records:
        month, asset_id = _to_month(record.get("Month")), record.get("Asset ID")
        line, amount = record.get("Revenue Line"), _num(record.get("Amount (₹Cr)"))
        if month is None or asset_id is None or line is None or amount is None:
            continue
        if str(asset_id) not in active and str(asset_id) != FUND_ID:
            report.warn(f"{SHEET_REVENUE_MIX}: row for unknown Asset ID {asset_id!r} ignored")
            continue
        result.append({"month": month, "assetId": str(asset_id), "line": str(line), "amount": amount})
    return result


def _read_tenant_sheet(wb, assets, report) -> tuple:
    """4_TENANT_STACK_LEASE holds three tables stacked vertically in the same columns."""
    rows = _rows(wb[SHEET_TENANT])
    profile_header = _header_row(rows, "Asset ID", SHEET_TENANT)
    expiry_marker = _find_row(rows, _first_cell_is("Lease Expiry Schedule"), profile_header + 1)
    stack_marker = _find_row(rows, _first_cell_is("Stack Plan"), profile_header + 1)
    markers = sorted(m for m in (expiry_marker, stack_marker) if m is not None)
    active = {a["id"] for a in assets}

    def next_marker(after):
        return next((m for m in markers if m > after), None)

    headers, records = _block(rows, profile_header, next_marker(profile_header))
    _require(headers, ("Asset ID", "Tenant Name", "Monthly Rent (₹Cr)", "Leasable Area (sft)", "Lease End Date"), SHEET_TENANT)
    tenants = []
    for record in records:
        name, asset_id = record.get("Tenant Name"), record.get("Asset ID")
        if name is None:
            continue
        if asset_id is None or str(asset_id) not in active:
            report.warn(f"{SHEET_TENANT}: tenant {name!r} has unknown Asset ID {asset_id!r}; ignored")
            report.dropped("tenant_profile")
            continue
        area = _num(record.get("Leasable Area (sft)"))
        tenants.append({
            "assetId": str(asset_id),
            "name": str(name),
            "monthlyRent": _num(record.get("Monthly Rent (₹Cr)")),
            "psf": _num(record.get("Current Rent PSF (INR)")),
            "area": area,
            "start": _to_date(record.get("Lease Start Date")),
            "end": _to_date(record.get("Lease End Date")),
            "notes": record.get("Notes"),
            "areaStatus": "OK" if area else "Needs area (sft)",
            "recordAsOf": _to_date(record.get("Record As Of")),
            "floor": record.get("Floor"),
            "category": record.get("Category"),
            "ownership": record.get("Ownership"),
        })
        if record.get("Lease End Date") is not None and tenants[-1]["end"] is None:
            report.warn(f"{SHEET_TENANT}[{asset_id}] {name}: Lease End Date {record.get('Lease End Date')!r} is not a date")

    schedule = []
    if expiry_marker is not None:
        _, records = _block(rows, expiry_marker + 1, next_marker(expiry_marker + 1))
        for record in records:
            asset_id, quarter = record.get("Asset ID"), record.get("Expiry Quarter")
            if asset_id is None or quarter is None:
                continue
            schedule.append({"assetId": str(asset_id), "quarter": str(quarter), "sft": _num(record.get("Area Expiring"))})

    stack = []
    if stack_marker is not None:
        headers, records = _block(rows, stack_marker + 1, next_marker(stack_marker + 1))
        tenant_column = next((h for h in ("Tenant / Occupier / Status", "Tenant Name") if h in headers), None)
        _require(headers, ("Asset ID", "Floor No.", "Area (K sft)", "Category"), f"{SHEET_TENANT} Stack Plan")
        for record in records:
            asset_id, floor = record.get("Asset ID"), record.get("Floor No.")
            occupant = record.get(tenant_column) if tenant_column else None
            if floor is None and occupant is None:
                continue   # preallocated row
            if asset_id is None or str(asset_id) not in active:
                report.warn(f"{SHEET_TENANT} Stack Plan: row {occupant!r} has unknown Asset ID {asset_id!r}; ignored")
                report.dropped("stack_plan")
                continue
            if isinstance(floor, float) and floor.is_integer():
                floor = int(floor)
            category = record.get("Category")
            ownership = record.get("Ownership")
            stack.append({
                "assetId": str(asset_id),
                "floor": floor if floor is not None else "—",
                "tenant": occupant,
                "areaK": _num(record.get("Area (K sft)")),
                "category": category.strip().lower() if isinstance(category, str) else None,
                "notes": record.get("Notes"),
                "ownership": ownership.strip().title() if isinstance(ownership, str) else "Fund",
            })
    return tenants, schedule, stack


def _read_registers(wb, assets, report) -> tuple:
    rows = _rows(wb[SHEET_REGISTERS])
    index = _header_row(rows, "Asset ID", SHEET_REGISTERS)
    header = [_canon(h) for h in rows[index]]
    split = next((i for i in range(1, len(header)) if header[i] == "Asset ID"), None)
    if split is None:
        raise ValueError(f"{SHEET_REGISTERS}: compliance register (second 'Asset ID' column) not found")
    active = {a["id"] for a in assets}

    def register(first, last, key, name):
        headers, records = _block(rows, index, None, first, last)
        _require(headers, ("Asset ID", key), f"{SHEET_REGISTERS} {name}")
        kept = []
        for record in records:
            asset_id = record.get("Asset ID")
            if record.get(key) is None:
                continue
            if asset_id is None or str(asset_id) not in active:
                report.warn(f"{SHEET_REGISTERS} {name}: {record.get(key)!r} has unknown Asset ID {asset_id!r}; ignored")
                report.dropped(name.lower())
                continue
            record["Asset ID"] = str(asset_id)
            kept.append(record)
        return kept

    approvals = register(0, split, "Approval Description", "Approvals")
    compliance = register(split, None, "Compliance Item", "Compliance")

    for name, records, key, status_col, allowed in (
        ("Approvals", approvals, "Approval Description", "Status", {"Complied", "In Progress", "Ongoing", "Revision Required"}),
        ("Compliance", compliance, "Compliance Item", "RAG Status", {"In Progress", "Ongoing", "Due Soon", "OK", "Not Due", "Not Applicable"}),
    ):
        seen, duplicates = set(), 0
        for record in records:
            identity = tuple(record.get(c) for c in ("Asset ID", key, "Due Date", "Validity"))
            duplicates += identity in seen
            seen.add(identity)
        if duplicates:
            report.warn(f"{SHEET_REGISTERS} {name}: {duplicates} row(s) repeat an earlier row for the same asset")
        unexpected = sorted({str(r.get(status_col)) for r in records if r.get(status_col) is not None} - allowed)
        blank = sum(1 for r in records if r.get(status_col) is None)
        if unexpected or blank:
            report.warn(f"{SHEET_REGISTERS} {name}: {blank} blank and {len(unexpected)} unexpected {status_col} value(s) {unexpected or ''}".rstrip())
    return approvals, compliance


def _check_tenants(tenants, schedule, assets, as_of, report):
    names = {a["id"]: a for a in assets}
    missing_area = [t for t in tenants if not t["area"]]
    if missing_area:
        report.warn(f"{SHEET_TENANT}: {len(missing_area)} tenant(s) have no Leasable Area, so they are left out of WALE and lease expiry")
    expired = [t for t in tenants if t["end"] and t["end"] < as_of]
    if expired:
        report.warn(f"{SHEET_TENANT}: {len(expired)} tenant lease(s) ended before {as_of:%d %b %Y}: " + ", ".join(f"{t['name']} ({t['assetId']})" for t in expired))

    derived = {}
    for tenant in tenants:
        if tenant["end"] and tenant["area"]:
            bucket = derived.setdefault(tenant["assetId"], {})
            label = quarter_label(tenant["end"])
            bucket[label] = bucket.get(label, 0) + tenant["area"]
    entered = {}
    for row in schedule:
        bucket = entered.setdefault(row["assetId"], {})
        bucket[row["quarter"]] = bucket.get(row["quarter"], 0) + (row["sft"] or 0)
        leasable = names.get(row["assetId"], {}).get("leasable")
        if leasable and row["sft"] and row["sft"] > leasable:
            report.warn(f"{SHEET_TENANT} Lease Expiry Schedule: {row['assetId']} {row['quarter']} shows {row['sft']:,.0f} sft, more than the asset's {leasable:,.0f} sft leasable area")
    stale = sorted(
        a for a in set(entered) | set(derived)
        if {q: round(v) for q, v in entered.get(a, {}).items()} != {q: round(v) for q, v in derived.get(a, {}).items()}
    )
    if schedule and stale:
        report.warn(
            f"{SHEET_TENANT}: the Lease Expiry Schedule block does not match tenant Lease End Dates for {', '.join(stale)}; "
            "the dashboard derives lease expiry from the tenant rows"
        )
    for asset in assets:
        if asset["workbookWale"] == 1.3 and any(t["assetId"] == asset["id"] and t["area"] and t["end"] for t in tenants):
            report.warn(
                f"{SHEET_ASSETS}[{asset['id']}]: WALE shows the formula's 1.3-year fallback (its range runs into the stacked "
                "expiry/stack tables); the dashboard recalculates WALE from tenant leases"
            )


# ── public entry point ────────────────────────────────────────────────────────

def ingest(filepath) -> FundWorkbook:
    filepath = Path(filepath)
    if not filepath.is_file():
        raise FileNotFoundError(f"Workbook not found: {filepath}")
    with warnings.catch_warnings():
        # openpyxl cannot read Excel's data-validation extension; it is irrelevant here.
        warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
        wb = openpyxl.load_workbook(str(filepath), data_only=True)

    missing = [name for name in REQUIRED_SHEETS if name not in wb.sheetnames]
    if missing:
        raise ValueError(f"Workbook missing required sheet(s): {missing}")

    report = IngestionReport()
    assets = _read_assets(wb, report)
    by_asset, months, report_month = _read_monthly(wb, assets, report)
    as_of = _month_end(report_month)
    fund = _roll_up_fund(by_asset, months)
    _check_fund_summary(wb, fund, report)

    monthly = {FUND_ID: fund, **by_asset}
    annual = {asset_id: _annual(series) for asset_id, series in monthly.items()}

    tenants, schedule, stack = _read_tenant_sheet(wb, assets, report)
    _check_tenants(tenants, schedule, assets, as_of, report)
    approvals, compliance = _read_registers(wb, assets, report)

    return FundWorkbook(
        meta=_read_meta(wb),
        assets=assets,
        months=months,
        report_month=report_month,
        as_of=as_of,
        monthly=monthly,
        annual=annual,
        rent_spread=_read_rent_spread(wb, assets, report),
        revenue_mix=_read_revenue_mix(wb, assets, report),
        tenants=tenants,
        lease_schedule=schedule,
        stack_plan=stack,
        approvals=approvals,
        compliance=compliance,
        data_quality=_read_data_quality(wb),
        report=report,
    )


if __name__ == "__main__":
    import sys

    default = Path(__file__).resolve().parents[1] / "data" / "real_estate_fund_data.xlsx"
    book = ingest(sys.argv[1] if len(sys.argv) > 1 else default)
    print(f"assets: {[a['id'] for a in book.assets]}")
    print(f"months: {[month_label(m) for m in book.months]} (report month {month_label(book.report_month)})")
    for row in book.monthly[FUND_ID]:
        print(
            f"  FUND {month_label(row['month'])}: leased {row['leased']:,.0f}/{row['leasable']:,.0f} sft "
            f"({row['occupancy']:.1f}%), rent ₹{row['monthlyRent']:.4f}Cr, collections ₹{row['collections']:.4f}Cr, "
            f"closing ₹{row['closingBalance']:.4f}Cr"
        )
    print(f"tenants: {len(book.tenants)}, stack rows: {len(book.stack_plan)}, "
          f"approvals: {len(book.approvals)}, compliance: {len(book.compliance)}")
    print(f"data quality: {book.data_quality['overall']}")
    for warning in book.report.warnings:
        print(f"[warn] {warning}")
